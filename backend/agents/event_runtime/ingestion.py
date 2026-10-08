import base64
from io import BytesIO
from zipfile import ZipFile
from hashlib import sha256
from defusedxml import ElementTree as ET
from pypdf import PdfReader
from docx import Document as WordDocument
from .models import Block, Document

MAX_BYTES = 10_000_000
MAX_TEXT = 2_000_000
MAX_BLOCK = 6000
OVERLAP = 600

class IngestionError(ValueError):
    pass

def unpack(document: Document, ocr=None) -> list[Block]:
    try:
        raw = base64.b64decode(document.content_base64, validate=True)
    except Exception as exc:
        raise IngestionError("invalid base64") from exc
    if not raw or len(raw) > MAX_BYTES:
        raise IngestionError("empty or oversized document")
    parts: list[tuple[str, str]] = []
    metadata = {}
    if document.media_type == "text/plain":
        parts = [("text", raw.decode("utf-8", errors="strict"))]
    elif document.media_type == "application/xml":
        # Generic XML is traversed, not interpreted as a validated E2B mapping.
        root = ET.fromstring(raw)
        def walk(node, path, depth=0):
            if depth > 100:
                raise IngestionError("XML depth limit")
            if node.text and node.text.strip():
                parts.append((path, node.text.strip()))
            # Coded E2B values need schema/OID-aware mapping: preserve attributes.
            if node.attrib:
                parts.append((path + "/@attributes", str(sorted(node.attrib.items()))))
            counts = {}
            for child in node:
                counts[child.tag] = counts.get(child.tag, 0) + 1
                walk(child, f"{path}/{child.tag}[{counts[child.tag]}]", depth + 1)
                if child.tail and child.tail.strip():
                    parts.append((path + "/tail", child.tail.strip()))
        walk(root, f"/{root.tag}[1]")
    elif document.media_type == "application/pdf":
        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted or len(reader.pages) > 1000:
            raise IngestionError("encrypted PDF or page limit")
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if list(page.images) or not text.strip():
                if ocr is None or i > 100 or float(page.mediabox.width)>2000 or float(page.mediabox.height)>2000:
                    raise IngestionError(f"page {i} requires supported OCR")
                text, info = ocr.pdf_page(raw, i)
                metadata[f"page:{i}"] = info
            parts.append((f"page:{i}", text))
    else:
        with ZipFile(BytesIO(raw)) as z:
            if sum(x.file_size for x in z.infolist()) > 50_000_000:
                raise IngestionError("DOCX expanded size limit")
            if any(x.filename.startswith("word/media/") for x in z.infolist()):
                raise IngestionError("DOCX contains images requiring visual/OCR inspection")
            if any(x.filename.startswith(("word/footnotes", "word/endnotes")) for x in z.infolist()):
                raise IngestionError("DOCX notes require an extended ingestion adapter")
        doc = WordDocument(BytesIO(raw))
        if doc.element.xpath(".//w:del | .//w:ins"):
            raise IngestionError("DOCX tracked changes require reviewed source version")
        # Use XML body order, retaining tables rather than dropping their text.
        for i, el in enumerate(doc.element.body):
            texts = el.xpath('.//w:t')
            if texts:
                parts.append((f"body:{i}", " ".join(t.text or "" for t in texts)))
        for section_index, section in enumerate(doc.sections):
            for kind in ("header", "footer", "first_page_header", "first_page_footer", "even_page_header", "even_page_footer"):
                region = getattr(section, kind)
                for ti, table in enumerate(region.tables):
                    parts.append((f"section:{section_index}/{kind}/table:{ti}", " ".join(c.text for r in table.rows for c in r.cells)))
                for i, para in enumerate(region.paragraphs):
                    if para.text.strip():
                        parts.append((f"section:{section_index}/{kind}:{i}", para.text))
    if sum(len(t) for _, t in parts) > MAX_TEXT:
        raise IngestionError("extracted text limit")
    blocks = []
    for locator, text in parts:
        if not text.strip():
            continue
        offset = 0
        while offset < len(text):
            end = min(offset + MAX_BLOCK, len(text))
            chunk = text[offset:end]
            bid = sha256(f"{document.id}:{locator}:{offset}:{chunk}".encode()).hexdigest()[:24]
            blocks.append(Block(id=bid, document_id=document.id, locator=f"{locator}/chars:{offset}-{end}", text=chunk, ingestion_metadata=metadata.get(locator, {})))
            if end == len(text):
                break
            offset = end - OVERLAP
    if not blocks:
        raise IngestionError("no readable text")
    return blocks
