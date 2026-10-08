import asyncio
import base64
import json
from io import BytesIO
from dataclasses import replace
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError
from backend.agents.event_runtime.models import Request, Document, Mention, Extraction, Verification, Finding
from backend.agents.event_runtime.ingestion import unpack, IngestionError, MAX_BLOCK
from backend.agents.event_runtime.engine import EventEngine
from backend.agents.event_runtime.api import create_app
from backend.agents.event_runtime.audit import AuditStore
from backend.agents.event_runtime.coding import MedDRADictionary


def doc(text, media_type='text/plain', id='d1'):
    raw = text.encode() if isinstance(text,str) else text
    return Document(id=id,media_type=media_type,language='en',content_base64=base64.b64encode(raw).decode())


def request(text='Patient P001 developed nausea.', **kwargs):
    return Request(case_id='c1',documents=[doc(text, **kwargs)])


class ScriptedProvider:
    """Test double only: fixed evidence used to exercise guards, not an AE model."""
    identity='test-double'
    def __init__(self, **changes):
        self.changes=changes
        self.calls=0
    async def extract(self, block):
        self.calls+=1
        start=block.text.find('nausea')
        if start < 0:
            return Extraction(mentions=[],unresolved=[])
        data=dict(block_id=block.id,start=start,end=start+6,verbatim='nausea',patient_id='P001',
                  patient_evidence='P001',assertion='affirmed',role='event',diagnosis_status='symptom_sign',
                  context_quote=block.text,rationale='explicit source mention',onset_quote=None,outcome_quote=None)
        data.update({k:v for k,v in self.changes.items() if k in Mention.model_fields})
        return Extraction(mentions=[Mention(**data)],unresolved=[])
    async def verify(self, block, extraction):
        if self.changes.get('fail'):
            raise RuntimeError('private clinical text must not escape')
        findings=[Finding(mention_index=i,disposition=self.changes.get('disposition','supported'),reason='test') for i in range(len(extraction.mentions))]
        if self.changes.get('missing_verifier'):
            findings=[]
        return Verification(findings=findings,missed_evidence=['nausea'] if self.changes.get('omission') else [])


def run(req=None, **changes):
    return asyncio.run(EventEngine(ScriptedProvider(**changes)).run(req or request()))


def test_supported_event_and_provenance():
    out=run()
    assert out.events[0].verbatim=='nausea'
    assert out.events[0].patient_id=='d1:P001'
    assert out.events[0].evidence[0].quote=='nausea'
    assert out.status=='review_required'
    assert out.qualified_for_autonomous_use is False
    assert out.provenance['document_sha256']['d1']
    assert out.coverage['clinical_completeness_certified'] is False


@pytest.mark.parametrize('changes',[
    {'start':0}, {'end':1000}, {'verbatim':'vomiting'}, {'block_id':'wrong'},
    {'context_quote':'invented context'}, {'onset_quote':'yesterday'},
    {'outcome_quote':'recovered'}, {'patient_evidence':'P002'},
    {'patient_id':'P002','patient_evidence':'P001'}])
def test_invalid_evidence_fails_closed(changes):
    out=run(**changes)
    assert out.events==[]
    assert out.status=='incomplete'
    assert out.observations[0]['disposition']=='invalid_evidence'


@pytest.mark.parametrize('assertion',['negated','historical','hypothetical'])
def test_nonaffirmed_context_retained(assertion):
    out=run(assertion=assertion)
    assert out.events==[]
    assert out.observations[0]['mention']['verbatim']=='nausea'


@pytest.mark.parametrize('role',['indication','medical_history','special_situation','outcome','aggregate','unknown'])
def test_role_separation(role):
    assert run(role=role).events==[]


def test_uncertain_event_retained():
    assert run(assertion='uncertain').events[0].assertion=='uncertain'


def test_unnamed_patient_preserved_for_review():
    out=run(patient_id=None,patient_evidence=None)
    assert out.events==[]
    assert out.observations[0]['disposition']=='unresolved_patient'


@pytest.mark.parametrize('disposition',['review','reject'])
def test_verifier_disagreement(disposition):
    assert run(disposition=disposition).events==[]


def test_verifier_contract_and_omission():
    assert run(missing_verifier=True).status=='incomplete'
    assert run(omission=True).status=='incomplete'


def test_provider_failure_not_empty_success():
    out=run(fail=True)
    assert out.status=='incomplete'
    assert out.coverage['failed_ids']
    assert 'private clinical' not in out.model_dump_json()


def test_empty_report_requires_review():
    assert run(request('Patient P001 reported no medical problems.')).status=='review_required'


def test_duplicate_document_ids_rejected():
    with pytest.raises(ValidationError):
        Request(case_id='x',documents=[doc('a'),doc('b')])


def test_invalid_base64_has_safe_result():
    d=doc('a'); d.content_base64='!'
    out=run(Request(case_id='c',documents=[d]))
    assert out.status=='incomplete'
    assert out.provenance['document_sha256']['d1'] is None


def test_multidocument_patient_not_merged():
    out=run(Request(case_id='x',documents=[doc('Patient P001 nausea.',id='a'),doc('Patient P001 nausea.',id='b')]))
    assert {e.patient_id for e in out.events}=={'a:P001','b:P001'}
    assert any('reconciliation' in x for x in out.issues)


def test_unicode_offsets():
    out=run(request('患者 P001: nausea.'))
    assert out.events[0].evidence[0].start==9


def test_chunking_preserves_every_character():
    source='x'*(MAX_BLOCK*3+17)
    blocks=unpack(doc(source))
    spans=[tuple(map(int,b.locator.rsplit('/chars:',1)[1].split('-'))) for b in blocks]
    assert spans[0][0]==0 and spans[-1][1]==len(source)
    assert all(a[1]>=b[0] for a,b in zip(spans,spans[1:]))
    assert all(len(b.text)<=MAX_BLOCK for b in blocks)


def test_repeated_episodes_not_collapsed():
    class Episodes(ScriptedProvider):
        async def extract(self,b):
            e=await super().extract(b)
            first=e.mentions[0]
            second=first.model_copy(update={'start':b.text.rfind('nausea'),'end':b.text.rfind('nausea')+6})
            return Extraction(mentions=[first,second],unresolved=[])
    out=asyncio.run(EventEngine(Episodes()).run(request('P001 nausea Monday; P001 nausea Friday.')))
    assert len(out.events)==2


def test_xml_paths_and_fail_closed_mapping():
    out=run(request('<case><patient>P001</patient><reaction>nausea</reaction></case>',media_type='application/xml'))
    assert out.status=='incomplete'
    assert any('ingestion_failed' in x for x in out.issues)
    assert unpack(doc('<a><b>x</b></a>',media_type='application/xml'))[0].locator.startswith('/a[1]/b[1]')


def test_xml_external_entity_rejected():
    value='<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>'
    assert run(request(value,media_type='application/xml')).status=='incomplete'


def test_blank_pdf_requires_ocr():
    from pypdf import PdfWriter
    writer=PdfWriter();writer.add_blank_page(width=200,height=200)
    out=BytesIO();writer.write(out)
    assert run(request(out.getvalue(),media_type='application/pdf')).status=='incomplete'


def test_docx_tables_preserved():
    from docx import Document as WD
    wd=WD();wd.add_paragraph('P001');t=wd.add_table(rows=1,cols=1);t.cell(0,0).text='nausea'
    out=BytesIO();wd.save(out)
    blocks=unpack(doc(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'))
    assert any('nausea' in b.text for b in blocks)


def test_dictionary_exact_versioned_only(tmp_path):
    path=tmp_path/'terms.json'
    path.write_text(json.dumps({'version':'TEST_ONLY','license_confirmed':True,'terms':[
        {'llt_code':'TEST1','llt_name':'nausea','pt_code':'TEST2','pt_name':'nausea','language':'en','current':True}]}))
    dictionary=MedDRADictionary(str(path))
    assert dictionary.suggest('Nausea','en')['version']=='TEST_ONLY'
    assert dictionary.suggest('nauseated','en') is None
    assert dictionary.suggest('nausea','fr') is None


def api(tmp_path):
    provider=ScriptedProvider()
    store=AuditStore(str(tmp_path/'audit.sqlite'))
    app=create_app(EventEngine(provider),store,{'a'*32:'tenant-a','b'*32:'tenant-b'})
    return app,provider,store


def headers(secret='a'*32,key='one'):
    return {'Authorization':'Bearer '+secret,'Idempotency-Key':key}


def test_api_auth_idempotency_conflict_and_tenant_isolation(tmp_path):
    app,provider,store=api(tmp_path)
    with TestClient(app) as client:
        assert client.post('/v1/events/extract',json=request().model_dump()).status_code==401
        a=client.post('/v1/events/extract',headers=headers(),json=request().model_dump())
        assert a.status_code==200
        again=client.post('/v1/events/extract',headers=headers(),json=request().model_dump())
        assert again.json()['run_id']==a.json()['run_id']
        assert provider.calls==1
        assert client.post('/v1/events/extract',headers=headers(),json=request('P001 nausea Friday').model_dump()).status_code==409
        b=client.post('/v1/events/extract',headers=headers('b'*32),json=request().model_dump())
        assert b.json()['run_id']!=a.json()['run_id']
        with store.connect() as db:
            assert db.execute('SELECT count(*) FROM audit').fetchone()[0]==2


def test_missing_idempotency_key(tmp_path):
    app,_,_=api(tmp_path)
    with TestClient(app) as client:
        assert client.post('/v1/events/extract',headers={'Authorization':'Bearer '+'a'*32},json=request().model_dump()).status_code==422


def test_audit_failure_blocks_response(tmp_path):
    app,_,store=api(tmp_path)
    def fail(*args): raise RuntimeError('database secret')
    store.save=fail
    with TestClient(app) as client:
        response=client.post('/v1/events/extract',headers=headers(),json=request().model_dump())
        assert response.status_code==503
        assert 'database secret' not in response.text


def test_api_body_limit(tmp_path):
    app,_,_=api(tmp_path)
    with TestClient(app) as client:
        assert client.post('/v1/events/extract',content=b'x'*15_000_001).status_code==413
